import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_design_document_docx(output_path, diagrams_dir):
    doc = docx.Document()

    # Set page margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # ----------------------------------------------------
    # Header / Title Block
    # ----------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("PhishGuard")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x0F, 0x1E, 0x36)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Design Document & System Architecture\nFunctional Block Diagram and UML Diagrams (Final Report)")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(13)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

    links_p = doc.add_paragraph()
    links_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links_p.paragraph_format.space_after = Pt(12)
    r1 = links_p.add_run("GitHub Repository: ")
    r1.font.bold = True
    links_p.add_run("https://github.com/ShreyanBangera00/PhishGuard\n")
    r2 = links_p.add_run("Live Deployment: ")
    r2.font.bold = True
    links_p.add_run("https://phish-guard-chi-taupe.vercel.app")

    # Author Table
    table = doc.add_table(rows=4, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Name", "Roll", "Sap Id"]
    data = [
        ["Shreyan Bangera", "K003", "70102400047"],
        ["Yuvraj Bishnoi", "K006", "70102400067"],
        ["Suraj Sahoo", "K001", "70102400074"]
    ]

    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(cell, "0F1E36")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, val in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = val
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=150, right=150)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    def add_h1(text):
        h = doc.add_heading(text, level=1)
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.runs[0].font.name = 'Arial'
        h.runs[0].font.size = Pt(14)
        h.runs[0].font.bold = True
        h.runs[0].font.color.rgb = RGBColor(0x0F, 0x1E, 0x36)
        return h

    def add_h2(text):
        h = doc.add_heading(text, level=2)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.runs[0].font.name = 'Arial'
        h.runs[0].font.size = Pt(12)
        h.runs[0].font.bold = True
        h.runs[0].font.color.rgb = RGBColor(0x14, 0x63, 0xF3)
        return h

    def add_image_with_caption(img_filename, caption_text, width_inches=6.0):
        full_path = os.path.join(diagrams_dir, img_filename)
        if not os.path.exists(full_path):
            print(f"WARNING: Image not found: {full_path}")
            return
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        run = p_img.add_run()
        run.add_picture(full_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Calibri'
        r_cap.font.italic = True
        r_cap.font.size = Pt(10)
        r_cap.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # ----------------------------------------------------
    # 1. Introduction
    # ----------------------------------------------------
    add_h1("1. Introduction")
    p = doc.add_paragraph(
        "PhishGuard is an explainable, multi-modal phishing detection platform designed to protect users against malicious URLs and deceptive social engineering emails. A user interacts with a modern corporate web dashboard to submit URLs for real-time lexical inspection or deep page scraping, or to analyze suspicious emails directly from their inbox."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph(
        "The FastAPI backend processes requests through dedicated feature extraction and threat intelligence layers. For URLs, it extracts 17 structural and lexical features, queries Google Safe Browsing API, fetches live DNS/WHOIS/TLS network intelligence, and evaluates risk with an ensemble Random Forest model accompanied by transparent, rule-based feature attribution explanations. For emails, the system evaluates social engineering tactics through the Adaptive Hybrid Multi-Factor Attribution (A-HMFA) engine, verifies sender authenticity via live DNS MX records and the Disify API, audits domain age with WHOIS, and offers direct 1-click email ingestion from Gmail via an SSL IMAP integration."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph(
        "This document describes the design and technical architecture of the PhishGuard platform using a functional block diagram and three types of UML diagrams: class diagrams, a state transition diagram, and sequence diagrams, fully updated to reflect the production source code in the repository."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(12)

    # ----------------------------------------------------
    # 2. Functional Block Diagram
    # ----------------------------------------------------
    add_h1("2. Functional Block Diagram")
    p = doc.add_paragraph(
        "Figure 1 shows the main functional blocks of PhishGuard and the data flow between them. The user interacts exclusively with the frontend dashboard. The dashboard communicates with the FastAPI backend over REST (JSON). The backend coordinates the feature acquisition layer, queries external threat intelligence and verification services (Google Safe Browsing, Disify Free Email API, live DNS and WHOIS), executes the analysis and ML layers, and synthesizes the final explainable response."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    # Insert Figure 1
    add_image_with_caption("page_2_img_1_Image15.png", "Figure 1: Functional block diagram of PhishGuard", width_inches=6.2)

    # Table 1: Modules and Responsibilities
    p = doc.add_paragraph("Table 1 lists each component with its location in the repository and its technical responsibilities.")
    p.paragraph_format.space_after = Pt(6)

    modules_data = [
        ["Frontend Dashboard", "frontend/\n(index.html, style.css, app.js)", "Interactive corporate UI; provides URL Scanner, Deep Content Scan, Email Analysis with A-HMFA visual graphs, Gmail IMAP inbox picker, and live network telemetry readouts."],
        ["FastAPI Backend", "src/api/main.py", "Exposes REST endpoints (/api/predict, /api/analyze-content, /api/analyze-email, /api/gmail/messages, /api/live-feed), validates schemas, coordinates service modules, and computes final hybrid risk scores."],
        ["URL Feature Extractor", "src/features/url_features.py", "Extracts 17 lexical, structural, and Shannon entropy features from raw URL strings in <2 ms."],
        ["Network Intelligence", "src/features/network_intel.py", "Gathers live host telemetry: DNS A-records, IP Geolocation / ASN (ip-api.com), WHOIS domain age and registrar (python-whois), and TLS certificate validity."],
        ["Live Sender Verification", "src/features/network_intel.py\n(verify_sender_live)", "Performs real-time sender verification: queries Disify Email API for live DNS MX mail server records and burner/disposable detection, combined with WHOIS domain age audits."],
        ["Email Analysis Engine", "src/features/email_features.py", "Implements the Adaptive Hybrid Multi-Factor Attribution (A-HMFA) model: zone-weighted logging (subject vs. body), credential theft detection, urgency scoring, cross-domain attack chain synergy (Ψ), and token density ML proxy (P_ML)."],
        ["Gmail IMAP Integration", "src/features/gmail_imap.py", "Connects to imap.gmail.com over SSL using Google App Passwords; extracts clean plain-text/HTML body, parses sender display names into claimed orgs, and enables 1-click scanning."],
        ["Page Analyzer (Deep Scan)", "src/scraper/page_analyzer.py", "Enforces SSRF validation (blocking private/loopback RFC-1918 IPs), fetches target DOM, detects password entry forms, brand impersonation mismatches, and external asset ratios."],
        ["Threat Intelligence Service", "src/threat_intel/threat_service.py", "Integrates Google Safe Browsing API v4 for authoritative real-time blacklisting of URLs and extracted email links."],
        ["ML Classifier Engine", "src/models/ml_engine.py,\nsrc/models/saved/", "Loads production Random Forest ensemble model to output calibrated phishing probabilities from tabular URL features."],
        ["Explainability Layer", "src/explainability/explainer.py", "Computes directional rule-based feature attribution scores, generating human-readable diagnostic factor cards (Critical, Warning, Safe)."],
        ["Offline Training Pipeline", "src/models/train.py,\ndata/prepare_dataset.py", "Prepares balanced 7,974 sample dataset, benchmarks Logistic Regression vs. Random Forest vs. GBDT on stratified 80/20 split, and exports champion artifacts."]
    ]

    t1 = doc.add_table(rows=len(modules_data) + 1, cols=3)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1_headers = ["Block / Module", "Repository Location", "Key Responsibilities"]

    for i, h in enumerate(t1_headers):
        cell = t1.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(cell, "0F1E36")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)

    for row_idx, r_data in enumerate(modules_data, start=1):
        for col_idx, text_val in enumerate(r_data):
            cell = t1.cell(row_idx, col_idx)
            cell.text = text_val
            if col_idx == 0:
                cell.paragraphs[0].runs[0].font.bold = True
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ----------------------------------------------------
    # 3. UML Diagrams
    # ----------------------------------------------------
    add_h1("3. UML Diagrams")

    # 3.1 Class Diagram
    add_h2("3.1 Class Diagram")
    p = doc.add_paragraph(
        "The class model is divided into two logical diagrams. Figure 2 illustrates the data models and the central FastAPI application interface. Request and response schemas inherit from Pydantic's BaseModel to guarantee strict data validation, serialization, and OpenAPI compliance. PhishGuardAPI encapsulates the REST routing endpoints, request handling, and response synthesis."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    # Insert Figure 2
    add_image_with_caption("page_4_img_1_Image20.png", "Figure 2: Class diagram - Data models and API interface", width_inches=5.8)

    p = doc.add_paragraph(
        "Figure 3 shows the service and engine components utilized by the PhishGuard backend. Python modules containing specialized functions are represented as stereotyped service classes, emphasizing loose coupling and modular separation of concerns. PhishGuardAPI coordinates these services depending on whether URL inspection, deep page scraping, or email analysis is invoked."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    # Insert Figure 3
    add_image_with_caption("page_5_img_1_Image23.png", "Figure 3: Class diagram - Analysis and feature services", width_inches=6.0)

    # 3.2 State Transition Diagram
    add_h2("3.2 State Transition Diagram")
    p = doc.add_paragraph(
        "Figure 4 illustrates the finite state lifecycle of an inspection request processed by the backend. A request transitions from Idle to Validating Input. Invalid payloads immediately transition to Rejected and return an HTTP 4xx response. Valid input progresses through parallel feature extraction, classification, explanation, and threat lookup states. If deep scraping is selected, the Page Analysis state verifies SSRF constraints and inspects DOM artifacts before entering the Risk Scoring state where hybrid heuristics and ML probabilities are synthesized."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    # Insert Figure 4
    add_image_with_caption("page_6_img_1_Image26.png", "Figure 4: State transition diagram of an analysis request", width_inches=5.6)

    # 3.3 Sequence Diagrams
    add_h2("3.3 Sequence Diagrams")
    p = doc.add_paragraph(
        "Figure 5 details the end-to-end execution flow for standard URL analysis (POST /api/predict). The FastAPI backend sequentially invokes the URL feature extractor, network intelligence module, machine learning classifier, explainability engine, and threat intelligence service before calculating the risk score and returning the PredictionResponse to the client."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    # Insert Figure 5
    add_image_with_caption("page_7_img_1_Image29.png", "Figure 5: Sequence diagram - URL analysis (POST /api/predict)", width_inches=5.8)

    p = doc.add_paragraph(
        "Figure 6 details the execution flow for deep content analysis (POST /api/analyze-content). In addition to lexical prediction, the Page Analyzer executes a strict Server-Side Request Forgery (SSRF) verification. If the URL resolves to loopback, broadcast, or private RFC-1918 networks, the request is halted with HTTP 400. Legitimate public endpoints are safely scraped to evaluate login forms, external assets, and SSL certificates."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    # Insert Figure 6
    add_image_with_caption("page_8_img_1_Image32.png", "Figure 6: Sequence diagram - Deep content analysis (POST /api/analyze-content)", width_inches=5.5)

    # ----------------------------------------------------
    # 4. Advanced Production Capabilities
    # ----------------------------------------------------
    add_h1("4. Advanced Capabilities & Pipeline Enhancements")
    p = doc.add_paragraph(
        "Through iterative software engineering and verification phases, PhishGuard incorporates several enterprise enhancements that extend beyond standard academic prototypes:"
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    enhancements = [
        ("Live Sender Verification (Disify API + WHOIS)",
         "Rather than relying on static sender address checks, the system performs live DNS MX mail server verification and disposable/temp-mail provider lookups via the Disify Free Email API. In parallel, python-whois extracts registration dates to flag newly created attack domains (<90 days old) vs. established enterprise domains."),
        ("Adaptive Hybrid Multi-Factor Attribution (A-HMFA v4)",
         "Email threat scoring implements continuous zone-weighted activation (counting subject line urgency hits at 2.0x weight vs. 1.0x in message body), evaluates cross-domain attack chain synergy (Ψ), and applies non-linear sigmoidal scaling to prevent premature saturation at 100."),
        ("Gmail IMAP Inbox Integration",
         "Users can directly connect their Gmail account over SSL IMAP using a 16-character Google App Password. The system fetches the latest emails from Inbox or Spam, converts complex MIME/HTML messages into clean readable text while preserving embedded links, automatically parses sender display names into claimed organizations, and runs 1-click threat audits."),
        ("Enterprise UI Design System",
         "The frontend dashboard adopts a professional corporate theme built with native HTML5, CSS3, and modern typography (Calibri/Arial, eliminating AI-like monospace fonts for numbers). It features interactive SVG risk score gauges, horizontal category progress bars, and transparent risk explanations."),
        ("Automated Test Suite",
         "The platform is continuously verified by 24 comprehensive pytest automated unit and integration tests covering the FastAPI endpoints, ML inference, SSRF security guards, TLS certificate extraction, Shannon entropy computation, and A-HMFA email scoring.")
    ]

    for title, desc in enhancements:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Inches(0.25)
        p_item.paragraph_format.space_after = Pt(4)
        r_t = p_item.add_run(f"• {title}: ")
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0x0F, 0x1E, 0x36)
        p_item.add_run(desc)

    doc.save(output_path)
    print(f"Successfully generated updated design document with embedded images at: {output_path}")

if __name__ == "__main__":
    current_dir = os.path.dirname(os.path.abspath(__file__))
    diagrams_dir = os.path.join(current_dir, "extracted_diagrams")
    out_docx = os.path.join(current_dir, "PhishGuard_Design_Document_Final_Report.docx")
    create_design_document_docx(out_docx, diagrams_dir)
